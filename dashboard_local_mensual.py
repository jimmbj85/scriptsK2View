import os
import sys
import subprocess
import importlib
import argparse
import asyncio

# ===========================================================================
# 0. AUTO-INSTALADOR INTELIGENTE (Script y EXE)
# ===========================================================================
def garantizar_dependencias():
    # Ruta del marcador para saber si Chromium ya fue instalado en el PC del usuario
    ruta_marcador = os.path.join(os.path.expanduser("~"), ".playwright_chromium_instalado")
    
    # Detectar si se está ejecutando como un ejecutable .exe compilado
    es_exe = getattr(sys, 'frozen', False)

    if es_exe:
        # MODO EXE: Las librerías ya están dentro, solo verificamos el navegador Chromium
        if not os.path.exists(ruta_marcador):
            print("====================================================================")
            print("⏳ CONFIGURACIÓN INICIAL: Descargando motor de gráficos...")
            print("Esto solo ocurrirá la primera vez. Por favor, no cierres la ventana.")
            print("====================================================================\n")
            try:
                from playwright.cli.main import main as playwright_main
                playwright_main(["install", "chromium"])
                with open(ruta_marcador, "w") as f:
                    f.write("instalado")
                print("\n✅ ¡Configuración completada con éxito! Iniciando el programa...\n")
            except Exception as e:
                print(f"\n❌ Error al configurar el motor de gráficos: {e}")
                input("Presiona Intro para salir...")
                sys.exit(1)
    else:
        # MODO SCRIPT: Comprobar e instalar librerías normales de Python si faltan
        librerias = ['pandas', 'openpyxl', 'playwright']
        necesita_instalacion = False
        for lib in librerias:
            try:
                importlib.import_module(lib)
            except ImportError:
                necesita_instalacion = True
                break

        if necesita_instalacion:
            print("====================================================================")
            print("⏳ CONFIGURACIÓN INICIAL: Instalando componentes necesarios...")
            print("====================================================================\n")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "pandas", "openpyxl", "playwright"])
                subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])
                print("\n✅ ¡Configuración completada con éxito!\n")
            except Exception as e:
                print(f"\n❌ Error en la configuración automática: {e}")
                sys.exit(1)

# Ejecutar la comprobación antes de cargar el resto de librerías pesadas
garantizar_dependencias()

# ===========================================================================
# IMPORTS DE LIBRERÍAS DE TRATAMIENTO DE DATOS
# ===========================================================================
import numpy as np
import pandas as pd
from playwright.async_api import async_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ===========================================================================
# CONFIGURACIÓN Y MAPAS DE DATOS
# ===========================================================================
ORDEN = ["1", "3", "5", "7"]
MAPA = {
    "1": {"n": "Búsqueda de datos",      "c": '#51a5e1', "hex_ox": "51A5E1", "i": 'fa-magnifying-glass'},
    "3": {"n": "Bajada de datos",        "c": '#5a6b7d', "hex_ox": "5A6B7D", "i": 'fa-cloud-arrow-down'},
    "5": {"n": "Otras consultas",        "c": '#23a985', "hex_ox": "23A985", "i": 'fa-comment-dots'},
    "7": {"n": "Cargas en Data Manager", "c": '#e75437', "hex_ox": "E75437", "i": 'fa-gear'}
}

MESES_MAP = [
    {"mes": "enero",      "nombre_real": "Enero",      "cod": "01"},
    {"mes": "febrero",    "nombre_real": "Febrero",    "cod": "02"},
    {"mes": "marzo",      "nombre_real": "Marzo",      "cod": "03"},
    {"mes": "abril",      "nombre_real": "Abril",      "cod": "04"},
    {"mes": "mayo",       "nombre_real": "Mayo",       "cod": "05"},
    {"mes": "junio",      "nombre_real": "Junio",      "cod": "06"},
    {"mes": "julio",      "nombre_real": "Julio",      "cod": "07"},
    {"mes": "agosto",     "nombre_real": "Agosto",     "cod": "08"},
    {"mes": "septiembre", "nombre_real": "Septiembre", "cod": "09"},
    {"mes": "octubre",    "nombre_real": "Octubre",    "cod": "10"},
    {"mes": "noviembre",  "nombre_real": "Noviembre",  "cod": "11"},
    {"mes": "diciembre",  "nombre_real": "Diciembre",  "cod": "12"}
]

async def generar_grafico(html_content, file_name):
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page(viewport={'width': 1080, 'height': 600})
        await page.set_content(html_content)
        await page.wait_for_timeout(1500)
        await page.screenshot(path=file_name)
        await browser.close()

def tiempo_natural_sin_findes(row):
    inicio = pd.to_datetime(row['Submit Date'])
    fin = pd.to_datetime(row['Completed Date'])
    if pd.isnull(inicio) or pd.isnull(fin) or fin < inicio: return 0
    segundos_totales = (fin - inicio).total_seconds()
    segundos_finde = 0
    current = inicio
    while current < fin:
        next_day = (current + pd.Timedelta(days=1)).normalize()
        if next_day > fin: next_day = fin
        if current.weekday() >= 5: segundos_finde += (next_day - current).total_seconds()
        current = next_day
    return (segundos_totales - segundos_finde) / 3600.0

# ===========================================================================
# LÓGICA PRINCIPAL
# ===========================================================================
async def main():
    if len(sys.argv) > 1:
        parser = argparse.ArgumentParser()
        parser.add_argument("--anio", type=str)
        parser.add_argument("--mes", type=str)
        args = parser.parse_args()
        anio_objetivo = args.anio
        mes_objetivo = args.mes.lower() if args.mes else ""
    else:
        print("====================================================")
        print("      ASISTENTE DE GENERACIÓN DE DASHBOARDS        ")
        print("====================================================\n")
        anio_objetivo = input("1. Introduce el año a procesar (ej. 2026): ").strip()
        mes_objetivo = input("2. Introduce el mes límite (ej. marzo): ").strip().lower()
        print("\n----------------------------------------------------")

    idx_limite = next((i for i, m in enumerate(MESES_MAP) if m["mes"] == mes_objetivo), None)
    if idx_limite is None:
        print(f"❌ El mes '{mes_objetivo}' no es válido. Asegúrate de escribir el nombre completo en español.")
        input("\nPresiona Intro para salir...")
        sys.exit(1)

    meses_a_procesar = MESES_MAP[:idx_limite + 1]
    print(f"🚀 Procesando datos desde Enero hasta {mes_objetivo.capitalize()} de {anio_objetivo}...\n")

    for config in meses_a_procesar:
        mes_inf, mes_cod = config['nombre_real'], config['cod']
        nombre_archivo = f"{anio_objetivo}_{mes_cod}_{mes_inf}_Informe Mensual Aprovisionamiento de Datos.xlsx"
        
        if not os.path.exists(nombre_archivo):
            print(f"⚠️ Archivo no encontrado (saltando): {nombre_archivo}")
            continue

        print(f"⏳ Procesando {mes_inf}...")
        file_img = f'dash_{mes_cod}.png'
        file_html = f'dash_{mes_cod}.html'

        try:
            xls = pd.ExcelFile(nombre_archivo)
            hojas = xls.sheet_names

            ws_9_name = next((h for h in hojas if h.startswith("9")), None)
            ratings_dict = {}
            if ws_9_name:
                df_9 = pd.read_excel(xls, sheet_name=ws_9_name)
                df_9.columns = [c.strip() for c in df_9.columns]
                if 'Categorization Tier 3' in df_9.columns:
                    df_9 = df_9.rename(columns={'Categorization Tier 3': 'Category'})
                if 'Rating' in df_9.columns and 'Category' in df_9.columns:
                    df_9['Category'] = df_9['Category'].str.strip()
                    replacements = {
                        "BUSQUEDA DE DATOS": "Búsqueda de datos", "OTRAS CONSULTAS": "Otras consultas",
                        "BAJADA DE DATOS": "Bajada de datos", "CARGAS EN DATA MANAGER": "Cargas en Data Manager"
                    }
                    df_9['Category'] = df_9['Category'].replace(replacements)
                    df_9['Rating'] = pd.to_numeric(df_9['Rating'], errors='coerce')
                    summary = df_9.groupby('Category')['Rating'].agg(['mean', 'count'])
                    ratings_dict = summary.to_dict('index')

            data_final = []
            total_peticiones = 0
            valores_sheets_peticiones = {}
            valores_sheets_horas = {}

            for id_p in ORDEN:
                ws_name = next((h for h in hojas if h.startswith(id_p)), None)
                if not ws_name: continue
                
                df = pd.read_excel(xls, sheet_name=ws_name)
                df.columns = [c.strip() for c in df.columns]
                df_flt = df[df['Status'].isin(["Terminado", "Cerrado"])].copy()

                for c in ['Submit Date', 'Completed Date']:
                    df_flt[c] = pd.to_datetime(df_flt[c].astype(str).str.replace(".","", regex=False), errors='coerce', format='mixed')

                df_flt['horas_res'] = df_flt.apply(tiempo_natural_sin_findes, axis=1)
                count = len(df_flt)
                media_horas = df_flt['horas_res'].mean() if count > 0 else 0
                
                nombre_cat = MAPA[id_p]['n']
                total_peticiones += count
                valores_sheets_peticiones[nombre_cat] = float(count)
                valores_sheets_horas[nombre_cat] = round(float(media_horas), 2)

                h_int, m_int = int(media_horas), int((media_horas % 1) * 60)
                stats = ratings_dict.get(nombre_cat, {'mean': 0, 'count': 0})
                rating_val = stats['mean'] if stats['count'] > 0 else 0
                count_val = int(stats['count']) if stats['count'] > 0 else 0
                rating_txt = f"{rating_val:.1f}/5 ({count_val})" if stats['count'] > 0 else "N/A"
                rating_bar_w = (rating_val * 20) if stats['count'] > 0 else 0

                data_final.append({
                    'cat': nombre_cat, 'val': count, 'media': media_horas, 'rating': float(rating_val),
                    'txt': f"{h_int:02d}h {m_int:02d}m", 'color': MAPA[id_p]['c'], 'icon': MAPA[id_p]['i'],
                    'r_txt': rating_txt, 'r_width': rating_bar_w, 'r_col': MAPA[id_p]['c'] if stats['count'] > 0 else "#e2e8f0"
                })

            wb = openpyxl.load_workbook(nombre_archivo)
            if "Dashboard_Resumen" in wb.sheetnames: del wb["Dashboard_Resumen"]
            
            ws_resumen = wb.create_sheet("Dashboard_Resumen", 0)
            ws_resumen.views.sheetView[0].showGridLines = True

            orden_columnas = ["Búsqueda de datos", "Bajada de datos", "Otras consultas", "Cargas en Data Manager"]
            ws_resumen.append(["Categoría"] + orden_columnas + ["TOTAL/MEDIA"])
            ws_resumen.append(["Tiempo (Horas)"] + [valores_sheets_horas.get(cat, 0) for cat in orden_columnas] + ['=ROUND(AVERAGE(B2:E2),2)'])
            ws_resumen.append(["Peticiones"] + [valores_sheets_peticiones.get(cat, 0) for cat in orden_columnas] + ['=SUM(B3:E3)'])
            ws_resumen.append(["Formato"] + [f'=INT({c}2) & " horas " & ROUND(({c}2-INT({c}2))*60, 0) & " minutos"' for c in ['B', 'C', 'D', 'E', 'F']])

            font_bold = Font(name="Arial", size=10, bold=True)
            font_white_bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            ws_resumen["A1"].font = font_white_bold
            ws_resumen["A1"].fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
            ws_resumen["A1"].alignment = Alignment(horizontal="left", vertical="center")

            for idx, id_p in enumerate(ORDEN):
                cell = ws_resumen.cell(row=1, column=idx + 2)
                cell.font = font_white_bold
                cell.fill = PatternFill(start_color=MAPA[id_p]['hex_ox'], end_color=MAPA[id_p]['hex_ox'], fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            ws_resumen["F1"].font = font_bold
            ws_resumen["F1"].alignment = Alignment(horizontal="center", vertical="center")

            for r in range(2, 5):
                ws_resumen.cell(row=r, column=1).font = font_bold
                ws_resumen.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
                for c in range(2, 7): ws_resumen.cell(row=r, column=c).alignment = Alignment(horizontal="right", vertical="center")

            for col in ws_resumen.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws_resumen.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(nombre_archivo)

            max_h = max([d['media'] for d in data_final]) or 1
            stops = []
            current_deg = 0
            for d in data_final:
                pct = (d['val'] / total_peticiones * 100) if total_peticiones > 0 else 0
                stops.append(f"{d['color']} {current_deg}% {current_deg + pct}%")
                current_deg += pct
            gradient = ",".join(stops)

            tiles_html = "".join([f'<div style="flex:1; height:65px; border:2px solid {d["color"]}; border-radius:8px; display:flex; flex-direction:column; align-items:center; justify-content:center; background:#fff;"><i class="fa-solid {d["icon"]}" style="font-size:16px; color:{d["color"]}; margin-bottom:4px;"></i><div style="font-weight:800; color:{d["color"]}; font-size:9px; letter-spacing:0.5px; text-align:center;">{d["cat"].upper()}</div></div>' for d in data_final])
            legend_grid = "".join([f'<div style="display: flex; align-items: center; gap: 6px;"><div style="width: 10px; height: 10px; background-color: {d["color"]}; border-radius: 2px;"></div><div style="font-size: 11px; color: {d["color"]}; font-weight: 700; white-space: nowrap;">{d["cat"]}: <b>{d["val"]}</b></div></div>' for d in data_final])
            barras_tiempo = "".join([f'<div style="display:flex; align-items:center; margin-bottom:12px;"><div style="width:160px; font-size:11px; font-weight:700; color:{d["color"]}; white-space:nowrap;">{d["cat"]}</div><div style="flex-grow:1; background:#f1f5f9; height:10px; border-radius:5px; overflow:hidden;"><div style="width:{int(d["media"]/max_h*100) if max_h>0 else 0}%; background:{d["color"]}; height:100%;"></div></div><div style="width:60px; text-align:right; font-size:11px; font-weight:700; color:{d["color"]};">{d["txt"]}</div></div>' for d in data_final])
            ratings_vertical = "".join([f'<div style="display: flex; align-items: center; margin-bottom: 12px; width: 100%;"><div style="width: 160px; font-size: 11px; font-weight: 800; color: {d["color"]}; text-transform: uppercase; text-align: right; padding-right: 15px; white-space: nowrap;">{d["cat"]}</div><div style="flex-grow: 1; background: #e2e8f0; height: 12px; border-radius: 6px; overflow: hidden;"><div style="width: {d["r_width"]}%; background: {d["color"]}; height: 100%;"></div></div><div style="width: 60px; text-align: left; font-size: 11px; font-weight: 700; color: {d["color"]}; padding-left: 10px;">{d["r_txt"]}</div></div>' for d in data_final])

            html_full = f"""<html><head><link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.5.1/css/all.min.css"></head>
            <body style="margin:0; padding:20px; background:#fff; width:1080px; height:600px; font-family:sans-serif; box-sizing:border-box;">
                <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:15px; border-bottom:2px solid #f1f5f9; padding-bottom:10px;">
                    <div style="font-weight:800; font-size:20px; color:#1e293b;">DASHBOARD OPERATIVO LOCAL</div>
                    <div style="font-weight:600; font-size:14px; color:#64748b; text-transform:uppercase;">{mes_inf} {anio_objetivo}</div>
                </div>
                <div style="display:flex; gap:10px; margin-bottom:15px;">{tiles_html}</div>
                <div style="display:flex; gap:20px; height:240px; margin-bottom: 20px;">
                    <div style="width:35%; border:1px solid #e2e8f0; border-radius:12px; padding:15px; display:flex; flex-direction:column; align-items:center; justify-content:space-between;">
                        <div style="font-weight:800; color:#475569; font-size:10px;">TOTAL PETICIONES</div>
                        <div style="width:110px; height:110px; border-radius:50%; background:conic-gradient({gradient}); display:flex; align-items:center; justify-content:center; position:relative;">
                            <div style="width:80px; height:80px; background:white; border-radius:50%; display:flex; align-items:center; justify-content:center; font-size:20px; font-weight:800; color:#1e293b;">{total_peticiones}</div>
                        </div>
                        <div style="display: grid; grid-template-columns: 1fr 1fr; gap: 8px; width: 100%;">{legend_grid}</div>
                    </div>
                    <div style="flex:1; border:1px solid #e2e8f0; border-radius:12px; padding:20px; display:flex; flex-direction:column; justify-content:center;">
                        <div style="font-weight:800; color:#475569; font-size:10px; margin-bottom:15px;">TIEMPO MEDIO RESOLUCIÓN</div>
                        {barras_tiempo}
                    </div>
                </div>
                <div style="width: 65%; margin: 0 auto; border:1px solid #e2e8f0; border-radius:12px; padding:20px; background:#f8fafc; display:flex; flex-direction:column; justify-content:center;">
                    <div style="font-weight:800; color:#475569; font-size:11px; margin-bottom:15px; text-align:center;">VALORACIÓN ENCUESTAS POR TIPO</div>
                    {ratings_vertical}
                </div>
            </body></html>"""
            
            with open(file_html, "w", encoding="utf-8") as f: f.write(html_full)
            await generar_grafico(html_full, file_img)
            print(f"   ↳ Realizado: Resumen en Excel, '{file_html}' y '{file_img}' creados.")

        except Exception as e:
            print(f"❌ Error procesando {mes_inf}: {e}")

    print("\n🏁 Proceso de automatización finalizado.")
    if len(sys.argv) == 1:
        input("\nPresiona Intro para cerrar esta ventana...")

if __name__ == "__main__":
    asyncio.run(main())